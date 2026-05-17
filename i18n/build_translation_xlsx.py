#!/usr/bin/env python3
"""Generate nisyros-wines-translation.xlsx — full translation worksheet (el).

Includes every translatable string on the site:
  • Static HTML strings (nav, intro slides, fixed UI labels, aria, SEO)
  • CMS — Config tab (hero, manifesto, contact, footer, wines page)
  • CMS — Wines tab (per-wine: name, subtitle, body, specs, detail fields)
  • CMS — Faces tab (per-photo: label, text)

The CMS rows are fetched live from the published Google Sheets CSV at build
time, so the workbook always reflects the current source of truth.

Re-run: `python3 i18n/build_translation_xlsx.py`
"""

import csv
import io
import sys
import urllib.request

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/home/user/nw/i18n/nisyros-wines-translation.xlsx"

CSV_BASE = ("https://docs.google.com/spreadsheets/d/e/"
            "2PACX-1vQGec_ewoWxtdcEXP05iJm4v2LHOoyW5sZc2bSBRVMzX7vlJIX8duf1JD"
            "--qMhpihBVgHMnHJxrgwkL/pub")
SHEET_GIDS = {
    "config": "1620319001",
    "wines":  "1993474932",
    "faces":  "1521558789",
}

HEADER_FILL = PatternFill("solid", fgColor="2B2725")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="F5F2EE")
SECTION_FILL = PatternFill("solid", fgColor="C46E4B")
SECTION_FONT = Font(name="Calibri", size=11, bold=True, color="F5F2EE", italic=True)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(border_style="thin", color="DDD6CC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def fetch_csv(name):
    url = f"{CSV_BASE}?gid={SHEET_GIDS[name]}&single=true&output=csv"
    with urllib.request.urlopen(url, timeout=30) as r:
        data = r.read().decode("utf-8")
    return list(csv.reader(io.StringIO(data)))


def write_header(ws, row, cols):
    for c, h in enumerate(cols, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def write_section_divider(ws, row, label, span):
    cell = ws.cell(row=row, column=1, value=label)
    cell.font = SECTION_FONT
    cell.fill = SECTION_FILL
    cell.alignment = Alignment(vertical="center")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)


def add_sheet(wb, name, cols, col_widths, rows, intro=None):
    """rows: list of dict (regular row) or {"__section__": "label"} (divider)."""
    ws = wb.create_sheet(name)
    r = 1
    if intro:
        ws.cell(row=r, column=1, value=intro).font = Font(italic=True, color="6B6360")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
        ws.row_dimensions[r].height = 32
        ws.cell(row=r, column=1).alignment = WRAP
        r += 2
    write_header(ws, r, cols)
    header_row = r
    r += 1
    for row in rows:
        if "__section__" in row:
            write_section_divider(ws, r, row["__section__"], len(cols))
            r += 1
            continue
        for c, key in enumerate(cols, 1):
            cell = ws.cell(row=r, column=c, value=row.get(key, ""))
            cell.alignment = WRAP
            cell.border = BORDER
        r += 1
    for c, h in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(c)].width = col_widths[c - 1]
    ws.freeze_panes = f"A{header_row + 1}"  # freeze header row
    return ws


# ══ Build workbook ════════════════════════════════════════════════════════
wb = Workbook()
wb.remove(wb.active)

# ── README ────────────────────────────────────────────────────────────────
readme = wb.create_sheet("README")
readme_lines = [
    ("Nisyros Wines — Greek translation workbook", "title"),
    ("", None),
    ("Purpose", "h2"),
    ("Translate every English string in column `en` into Greek in column `el`.", None),
    ("Leave column `key` and `en` untouched — they wire the translation back into the site.", None),
    ("", None),
    ("How it's organized", "h2"),
    ("Each sheet groups related strings. Order follows the user's path through the site:", None),
    ("  1. README              — this page", None),
    ("  2. Static HTML         — nav, intro slides, fixed UI labels (Story/Winemaking/…), aria labels", None),
    ("  3. CMS — Config        — hero, full manifesto, contact texts, footer, collection descriptions", None),
    ("  4. CMS — Wines         — per-wine content (name, subtitle, body, specs, detail fields) × 5 wines", None),
    ("  5. CMS — Faces         — gallery photo captions (label + text) × ~27 photos", None),
    ("  6. SEO                 — page title, meta description, social tags, structured data", None),
    ("", None),
    ("Guidelines", "h2"),
    ("• Tone: short declarative sentences, manifesto-like, no marketing fluff.", None),
    ("• Multi-line titles split with slashes (\"Wines / from the / Edge\") render one word per line — keep that rhythm in Greek.", None),
    ("• Brand names stay in latin script: Nisyros Wines, Monopàtia, Nereides, Volcanica, 3,2,1, 40 Milia Konda, Roudià, Apiri, Atmida.", None),
    ("• Grape varieties: keep latin (Athiri, Mandilaria, Assyrtiko, Mavrothiriko) unless the brand prefers Greek script — confirm with Nicola.", None),
    ("• Geographic names: use the natural Greek form — Νίσυρος, Δωδεκάνησα, Αιγαίο, Ρόδος, Έμπωνας, Σιάννα.", None),
    ("• Numbers, percentages, volumes, alcohol degrees, altitudes stay as-is.", None),
    ("• Body text with `||` separators: keep the `||` exactly where it is. Each part is rendered in its own block (Story / Winemaking / Tasting Notes).", None),
    ("• Specs separated by `|`: keep the `|` separators between items.", None),
    ("• If a string contains `\\n`, that is a literal newline marker — keep it.", None),
    ("• If meaning is ambiguous, write a question in the `notes` column.", None),
    ("", None),
    ("Deliverable", "h2"),
    ("Save this file as-is when done. The developer copies `el` values back into the codebase and into the Google Sheet's new `_el` columns.", None),
    ("", None),
    ("Generated from the live Google Sheet — re-run i18n/build_translation_xlsx.py to refresh.", "footer"),
]
for i, (text, kind) in enumerate(readme_lines, 1):
    cell = readme.cell(row=i, column=1, value=text)
    if kind == "title":
        cell.font = Font(size=18, bold=True, color="C46E4B")
    elif kind == "h2":
        cell.font = Font(size=13, bold=True, color="2B2725")
    elif kind == "footer":
        cell.font = Font(size=10, italic=True, color="6B6360")
    else:
        cell.font = Font(size=11, color="2B2725")
    cell.alignment = Alignment(wrap_text=True, vertical="top")
readme.column_dimensions["A"].width = 120

# ── 2. STATIC HTML ───────────────────────────────────────────────────────
STATIC_COLS = ["key", "en", "el", "notes"]
STATIC_W = [32, 60, 60, 50]

static_rows = [
    {"__section__": "NAV — top navigation"},
    {"key": "nav.manifesto", "en": "Manifesto",  "notes": "Top nav link"},
    {"key": "nav.wines",     "en": "Wines",      "notes": "Top nav link"},
    {"key": "nav.faces",     "en": "Faces",      "notes": "Top nav link"},
    {"key": "nav.contact",   "en": "Contact",    "notes": "Top nav link"},
    {"key": "nav.lang_en",   "en": "EN",         "notes": "Language switcher label (English)"},
    {"key": "nav.lang_el",   "en": "EL",         "notes": "Language switcher label (Greek) — likely 'ΕΛ'"},
    {"key": "nav.aria_menu", "en": "Menu",       "notes": "Hamburger button aria-label (screen reader)"},

    {"__section__": "WINES — intro slide (above the bottle cards)"},
    {"key": "wines.intro.count",       "en": "5 Wines.",     "notes": "Big line 1, red"},
    {"key": "wines.intro.vintages",    "en": "2 Vintages.",  "notes": "Big line 2"},
    {"key": "wines.intro.collections", "en": "3 Collections.", "notes": "Big line 3"},
    {"key": "wines.intro.cta",         "en": "Discover >",   "notes": "CTA button"},

    {"__section__": "WINES — card body labels (rendered above each text block)"},
    {"key": "wines.label.story",         "en": "Story",        "notes": "Wine card — label above part 1 of body"},
    {"key": "wines.label.winemaking",    "en": "Winemaking",   "notes": "Wine card — label above part 2 of body"},
    {"key": "wines.label.tasting_notes", "en": "Tasting Notes","notes": "Wine card — label above part 3 of body"},

    {"__section__": "WINES — detail row labels (right column of each card)"},
    {"key": "wines.detail.vintages",     "en": "Vintages",     "notes": "Wine detail row label"},
    {"key": "wines.detail.type",         "en": "Type",         "notes": "Wine detail row label"},
    {"key": "wines.detail.variety",      "en": "Variety",      "notes": "Wine detail row label"},
    {"key": "wines.detail.appellation",  "en": "Appellation",  "notes": "Wine detail row label"},
    {"key": "wines.detail.origin",       "en": "Origin",       "notes": "Wine detail row label"},
    {"key": "wines.detail.altitude",     "en": "Altitude",     "notes": "Wine detail row label"},
    {"key": "wines.detail.soil",         "en": "Soil",         "notes": "Wine detail row label"},
    {"key": "wines.detail.vinification", "en": "Vinification", "notes": "Wine detail row label"},
    {"key": "wines.detail.production",   "en": "Production",   "notes": "Wine detail row label"},
    {"key": "wines.collection_suffix",   "en": "Collection",
     "notes": 'Appended after name (e.g. "Volcanica Collection"). In Greek likely "Συλλογή" placed BEFORE the name — flag if word order matters.'},
    {"key": "wines.alc_template",        "en": "alc.{n}% by vol.",
     "notes": "Wine card footer template. Keep the {n} placeholder."},

    {"__section__": "WINES — collections slide (final card)"},
    {"key": "wines.collections.eyebrow", "en": "Our Collections", "notes": "Small eyebrow over title"},
    {"key": "wines.collections.title",   "en": "Three Directions","notes": "Big title — collection descriptions live in CMS Config"},

    {"__section__": "FACES — section header + lightbox"},
    {"key": "faces.title.line1", "en": "Faces",    "notes": "Section title line 1"},
    {"key": "faces.title.line2", "en": "& Places", "notes": "Section title line 2"},
    {"key": "faces.intro",
     "en": "The island. The people. The bottles. What the cellar smells like at five in the morning.",
     "notes": "Section intro caption"},
    {"key": "faces.aria.close",    "en": "Close",    "notes": "Lightbox close button — aria label"},
    {"key": "faces.aria.previous", "en": "Previous", "notes": "Lightbox previous button — aria label"},
    {"key": "faces.aria.next",     "en": "Next",     "notes": "Lightbox next button — aria label"},

    {"__section__": "CONTACT — section title + block labels (block contents are in CMS Config)"},
    {"key": "contact.title.line1", "en": "Find", "notes": "Big section title — line 1"},
    {"key": "contact.title.line2", "en": "Us",   "notes": "Big section title — line 2"},
    {"key": "contact.label.location", "en": "Location",          "notes": "Block label"},
    {"key": "contact.label.contact",  "en": "Contact",           "notes": "Block label"},
    {"key": "contact.label.visits",   "en": "Visits & Tastings", "notes": "Block label"},
    {"key": "contact.label.follow",   "en": "Follow",            "notes": "Block label"},
    {"key": "contact.label.trade",    "en": "Trade & Distribution", "notes": "Block label"},
    {"key": "contact.map_link",       "en": "See on map",        "notes": "Google Maps link"},
    {"key": "contact.email_link",     "en": "e-mail",            "notes": "Mailto link text"},
    {"key": "contact.whatsapp_link",  "en": "whatsapp",          "notes": "WhatsApp link text"},
    {"key": "contact.social_note",
     "en": "The vineyard, the cellar, the island. Unfiltered.",
     "notes": "Sub-note under @nisyroswines handle"},
]
add_sheet(wb, "2. Static HTML", STATIC_COLS, STATIC_W, static_rows,
          intro="Strings hard-coded in sections/*.html and js/main.js. These are NOT in the CMS — they live in the code and must be translated here.")

# ── 3. CMS — Config ──────────────────────────────────────────────────────
config_rows_raw = fetch_csv("config")
CFG_NOTES = {
    "hero_location":      "Caption above the wordmark on the hero",
    "hero_tagline":       "Currently unused on site — translate anyway in case it's wired up later",
    "m1_line1":  "Manifesto I — Big title line 1",
    "m1_line2":  "Manifesto I — Big title line 2",
    "m1_line3":  "Manifesto I — Big title line 3 (red accent)",
    "m1_body":   "Manifesto I — Body paragraph",
    "m2_line1":  "Manifesto II — Big title line 1",
    "m2_line2":  "Manifesto II — Big title line 2 (red accent)",
    "m2_line3":  "Manifesto II — Empty in current data; leave empty",
    "m2_body":   "Manifesto II — Body paragraph",
    "m3_line1":  "Manifesto III — Big title line 1",
    "m3_line2":  "Manifesto III — Big title line 2 (red accent)",
    "m3_body":   "Manifesto III — Body paragraph",
    "m4_line1":  "Manifesto IV — Big title line 1",
    "m4_line2":  "Manifesto IV — Big title line 2",
    "m4_line3":  "Manifesto IV — Big title line 3 (red accent)",
    "m4_body":   "Manifesto IV — Body paragraph",
    "m5_line1":  "Manifesto V — Big title line 1",
    "m5_line2":  "Manifesto V — Big title line 2",
    "m5_line3":  "Manifesto V — Big title line 3 (red accent)",
    "m5_body":   "Manifesto V — Body paragraph",
    "m6_line1":  "Manifesto VI — Big title line 1",
    "m6_line2":  "Manifesto VI — Big title line 2 (red accent)",
    "m6_body":   "Manifesto VI — Body paragraph. Contains a literal newline.",
    "credo_line1": "Credo / Manifesto VII — Big title line 1",
    "credo_line2": "Credo / Manifesto VII — Big title line 2",
    "credo_line3": "Credo / Manifesto VII — Big title line 3 (red accent)",
    "credo_body":  "Credo / Manifesto VII — Body paragraph",
    "contact_subtitle":     "Sub-title under \"Find Us\"",
    "contact_loc_value":    "Big value in Location block — \"Nisyros\" → \"Νίσυρος\"",
    "contact_loc_sub":      "Address. `\\n` = line break, keep them.",
    "contact_email_value":  "Email address — DO NOT TRANSLATE (it's the literal mailto target)",
    "contact_email_sub":    "Sub-note under contact links",
    "contact_visits_value": "Big value in Visits block",
    "contact_visits_sub":   "Sub-note in Visits block",
    "contact_trade_value":  "Big value in Trade block",
    "contact_trade_sub":    "Sub-note in Trade block",
    "footer_text":          "Page footer. Keep the · separators.",
    "wines_subtitle":       "Wines page subtitle (not yet wired into HTML — translate anyway)",
    "coll_1_name":          "Collection name — likely stays \"Monopàtia\"",
    "coll_1_desc":          "Monopàtia description (also appears on Three Directions slide)",
    "coll_2_name":          "Collection name — likely stays \"Nereides\"",
    "coll_2_desc":          "Nereides description (also appears on Three Directions slide)",
    "coll_3_name":          "Collection name — likely stays \"Volcanica\"",
    "coll_3_desc":          "Volcanica description (also appears on Three Directions slide)",
}
config_rows = []
for row in config_rows_raw:
    if not row or len(row) < 2:
        continue
    key, val = row[0], row[1]
    if not key.strip():
        continue
    if key.startswith("▸"):
        config_rows.append({"__section__": key.strip().lstrip("▸").strip()})
        continue
    config_rows.append({
        "key":   key,
        "en":    val,
        "notes": CFG_NOTES.get(key, ""),
    })
add_sheet(wb, "3. CMS — Config", STATIC_COLS, STATIC_W, config_rows,
          intro="Lives in Google Sheets (tab `config`). These values overwrite the HTML at runtime — translate them all. Section dividers are for orientation only.")

# ── 4. CMS — Wines ───────────────────────────────────────────────────────
wines_rows_raw = fetch_csv("wines")

# Header row in the CSV is row index 2 (after the two banner rows)
header = [c.split("\n")[0].strip().lower() for c in wines_rows_raw[2]]


def w_get(row, col):
    try:
        idx = header.index(col.lower())
    except ValueError:
        return ""
    return row[idx] if idx < len(row) else ""


WINE_COLS = ["wine", "field", "en", "el", "notes"]
WINE_W = [16, 18, 60, 60, 38]

WINE_FIELDS = [
    ("name",         "Wine name (likely stays in original)"),
    ("subtitle",     "One-liner under the wine name"),
    ("type",         "Wine type — e.g. White / Rosé / Light Red / Orange"),
    ("body",         "3 parts separated by ||  →  Story || Winemaking || Tasting Notes. KEEP the || separators."),
    ("specs",        "Pipe-separated specs shown as small tags. KEEP the | separators between items."),
    ("variety",      "Grape varieties — keep latin names (Athiri, Mandilaria…) unless the brand prefers Greek script"),
    ("appellation",  "Designation — e.g. PGI Dodecanese"),
    ("origin",       "Vineyard / location"),
    ("altitude",     "Numbers usually stay (e.g. 400-900m)"),
    ("soil",         "Soil description"),
    ("vinification", "Brief winemaking note"),
    ("production",   "Production count — number stays"),
    ("vintages",     "Years separated by `|` — keep numbers and separator"),
]
wines_rows = []
# CSV format: row 0 = banner, row 1 = instructions, row 2 = header,
# rows 3+ = either section dividers (▸…) or wines.
for row in wines_rows_raw[3:]:
    if not row or not any(c.strip() for c in row):
        continue
    first = row[0].strip()
    if first.startswith("▸"):
        wines_rows.append({"__section__": first.lstrip("▸").strip()})
        continue
    if first.startswith("📷") or first.startswith("HOW"):
        continue
    wine_id = w_get(row, "id")
    name = w_get(row, "name")
    if not wine_id or not name:
        continue
    label = f"{wine_id} — {name}"
    for field, note in WINE_FIELDS:
        val = w_get(row, field).strip()
        if not val:
            continue
        wines_rows.append({
            "wine":  label,
            "field": field,
            "en":    val,
            "el":    "",
            "notes": note,
        })
add_sheet(wb, "4. CMS — Wines", WINE_COLS, WINE_W, wines_rows,
          intro="Lives in Google Sheets (tab `wines`). One row per (wine × translatable field). Final values will be written back into new `_el` columns in the sheet.")

# ── 5. CMS — Faces ───────────────────────────────────────────────────────
faces_rows_raw = fetch_csv("faces")
# CSV: row 0 = banner, row 1 = instructions, row 2 = header, rows 3+ = photos
FACES_COLS = ["photo", "field", "en", "el", "notes"]
FACES_W = [10, 12, 50, 50, 40]
faces_rows = []
for row in faces_rows_raw[3:]:
    if not row or len(row) < 4:
        continue
    order, _img, label, text = (row + [""] * 4)[:4]
    order = order.strip()
    if not order:
        continue
    if order.startswith("📷") or order.startswith("HOW"):
        continue
    if label.strip():
        faces_rows.append({
            "photo": f"#{order}",
            "field": "label",
            "en":    label,
            "el":    "",
            "notes": "Short caption shown over the photo in the lightbox (e.g. Landscape, Winery, Harvest, People). Some labels repeat across photos.",
        })
    if text.strip():
        faces_rows.append({
            "photo": f"#{order}",
            "field": "text",
            "en":    text,
            "el":    "",
            "notes": "One-line caption under the label.",
        })
add_sheet(wb, "5. CMS — Faces", FACES_COLS, FACES_W, faces_rows,
          intro="Lives in Google Sheets (tab `faces`). One row per (photo × translatable field). `#N` matches the `oredr` column in the sheet. Final values will be written back into new `_el` columns.")

# ── 6. SEO ───────────────────────────────────────────────────────────────
seo_rows = [
    {"key": "seo.title",
     "en": "Nisyros Wines — Volcanic Natural Wines from the Aegean Sea",
     "notes": "<title> tag — try to stay under ~65 chars in Greek"},
    {"key": "seo.description",
     "en": "Small producer of natural wines on Nisyros, a volcanic island in the Dodecanese, Greece. Indigenous grape varieties, minimal intervention, volcanic terroir.",
     "notes": "<meta description>. A draft Greek version already exists in head.html — review and refine: \"Μικρός παραγωγός φυσικών κρασιών στη Νίσυρο, ηφαιστειογενές νησί στα Δωδεκάνησα. Αυτόχθονες ποικιλίες, ελάχιστη παρέμβαση, ηφαιστειακό terroir.\""},
    {"key": "seo.keywords",
     "en": "natural wines Greece, volcanic wines, Aegean wines, Nisyros, Dodecanese wines, Greek natural wine producer",
     "notes": "<meta keywords>. Greek draft already in head.html: \"φυσικά κρασιά Ελλάδα, ηφαιστειακά κρασιά, κρασιά Αιγαίου, Νίσυρος, Δωδεκάνησα, Ελληνικά φυσικά κρασιά\""},
    {"key": "seo.og.title",
     "en": "Nisyros Wines — Volcanic Natural Wines from the Aegean Sea",
     "notes": "Open Graph title (social share preview). Usually = seo.title."},
    {"key": "seo.og.description",
     "en": "Small producer of natural wines on Nisyros, a volcanic island in the Dodecanese, Greece.",
     "notes": "Open Graph description (shorter than seo.description)."},
    {"key": "seo.og.locale", "en": "en_GB",
     "notes": "For Greek version use `el_GR` — fixed value, no translation."},
    {"key": "seo.schema.description",
     "en": "Small producer of natural wines on Nisyros, a volcanic island in the Dodecanese, Greece. Indigenous grape varieties, minimal intervention, volcanic terroir.",
     "notes": "Schema.org Winery `description`. Usually = seo.description."},
]
add_sheet(wb, "6. SEO", STATIC_COLS, STATIC_W, seo_rows,
          intro="Meta tags and structured data. A draft Greek version of description/keywords already exists in head.html — review and refine, don't rewrite from scratch.")

wb.save(OUT)
print("✓ Wrote", OUT)
print("  Sheets:", ", ".join(s for s in wb.sheetnames))
